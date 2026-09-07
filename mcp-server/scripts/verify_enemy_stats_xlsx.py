# Compares `helldivers.raw_enemy_stats` (populated by db/import_raw_sheet.cjs from the live
# DiversDex 2.3 gviz CSV export) against a manually-exported xlsx of the same "Enemy stats" tab,
# cell by cell (percent-formatted columns are normalized numerically so "30%" == 0.3 == "30%").
# Writes any true gaps (xlsx has a value, DB doesn't) to
# mcp-server/cache/enemy_stats_xlsx_patch.json for db/patch_raw_enemy_stats.cjs to apply.
# Update XLSX_PATH below to point at whatever export you were just given.
XLSX_PATH = "D:/dev_study/hldv/enemystat.xlsx"

import openpyxl, datetime, os, re, sys, json
import psycopg2

env = {}
with open("D:/dev/helldivers2dex/db/.env", "r", encoding="utf-8") as f:
    for line in f:
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        env[k.strip()] = v.strip().strip('"').strip("'")

DATABASE_URL = env.get("DATABASE_URL")
if not DATABASE_URL:
    print("No DATABASE_URL found in db/.env")
    sys.exit(1)

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb["Sheet1"]
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

def sanitize_col(name, index, seen):
    s = (name or "").lower()
    s = s.replace("\n", " ")
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = s.strip("_")
    if not s:
        s = "col_" + str(index)
    out = s
    i = 2
    while out in seen:
        out = s + "_" + str(i)
        i += 1
    seen.add(out)
    return out

seen = set()
db_cols = [sanitize_col(h, i, seen) for i, h in enumerate(headers)]

def cell_to_text(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return f"{v.month}/{v.day}"
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return str(v)
    return str(v).strip()

def norm(v):
    v = (v or "").strip()
    if v in ("--", "-", ""):
        return ""
    is_pct = v.endswith("%")
    if is_pct:
        v = v[:-1]
    try:
        f = float(v)
        if is_pct:
            f = f / 100.0
        f = round(f, 6)
        if f == int(f):
            return str(int(f))
        return str(f)
    except ValueError:
        return v

xlsx_rows = []
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if all(v is None or str(v).strip() == "" for v in vals):
        continue
    xlsx_rows.append([cell_to_text(v) for v in vals])

print(f"xlsx data rows: {len(xlsx_rows)}, columns: {len(db_cols)}")

conn = psycopg2.connect(DATABASE_URL)
conn.autocommit = True
cur = conn.cursor()
cur.execute("SET search_path TO helldivers")
col_list = ", ".join('"' + c + '"' for c in db_cols)
cur.execute(f"SELECT row_num, {col_list} FROM raw_enemy_stats ORDER BY row_num")
db_rows = cur.fetchall()
print(f"db rows: {len(db_rows)}")

mismatches = []
if len(db_rows) != len(xlsx_rows):
    print(f"!! ROW COUNT MISMATCH: xlsx={len(xlsx_rows)} db={len(db_rows)}")

for i in range(min(len(xlsx_rows), len(db_rows))):
    db_row = db_rows[i][1:]
    xl_row = xlsx_rows[i]
    for j, col in enumerate(db_cols):
        a = norm(xl_row[j] if j < len(xl_row) else "")
        b = norm(db_row[j] if j < len(db_row) else "")
        if a != b:
            mismatches.append((i, col, xl_row[j] if j < len(xl_row) else None, db_row[j] if j < len(db_row) else None))

print(f"mismatches: {len(mismatches)}")
real_conflicts = [m for m in mismatches if norm(m[2]) and norm(m[3])]
fills = [m for m in mismatches if not norm(m[3]) and norm(m[2])]
db_only = [m for m in mismatches if norm(m[3]) and not norm(m[2])]
print(f"real_conflicts (both non-empty, different): {len(real_conflicts)}")
for m in real_conflicts:
    print("CONFLICT", m)
print(f"fills (xlsx has value, db empty): {len(fills)}")
for m in fills[:60]:
    print("FILL", m)
print(f"db_only (db has value, xlsx empty): {len(db_only)}")
for m in db_only[:30]:
    print("DB_ONLY", m)

patch = [{"row_num": m[0], "col": m[1], "value": m[2]} for m in fills]
with open("D:/dev/helldivers2dex/mcp-server/cache/enemy_stats_xlsx_patch.json", "w", encoding="utf-8") as f:
    json.dump(patch, f, ensure_ascii=False, indent=2)
print(f"wrote patch with {len(patch)} cells")

cur.close()
conn.close()
